#!/usr/bin/env python3
"""Task C live deck (16:9). Offline: reads saved charts and evidence only.

    python deck.py
"""
from __future__ import annotations

import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
TC = ROOT / "taskC_assessment"
RCH = TC / "charts"
EV = TC / "evidence"
TA = ROOT / "taskA_holidays" / "charts"
DCH = HERE / "charts"
OUT = HERE / "TaskC_Live_Presentation.pptx"


def rgb(h: str) -> RGBColor:
    h = h.lstrip("#")
    return RGBColor(int(h[:2], 16), int(h[2:4], 16), int(h[4:], 16))


NAVY, TEAL, GREEN = rgb("#14213D"), rgb("#0F8B8D"), rgb("#2A9D8F")
AMBER, CORAL, BLUE = rgb("#E9A23B"), rgb("#E76F51"), rgb("#3A86FF")
SLATE, INK, WHITE = rgb("#5C677D"), rgb("#1B1B1E"), rgb("#FFFFFF")
PAPER, LINE, MIST = rgb("#FAFBFD"), rgb("#DDE2EA"), rgb("#EEF2F7")

SECTIONS = [
    ("Context", SLATE),
    ("Task A calendar", AMBER),
    ("Data consistency", TEAL),
    ("Behaviour under load", CORAL),
    ("Risk, incident, improvements", GREEN),
]
STATUS = {
    "verified": ("VERIFIED", GREEN),
    "partial": ("VERIFIED · CAUSE OPEN", AMBER),
    "hypothesis": ("HYPOTHESIS", AMBER),
    "corrected": ("VERIFIED · CORRECTED", AMBER),
    "not done": ("NOT DONE", CORAL),
    "proposal": ("PROPOSAL", BLUE),
    "simulation": ("SIMULATION", SLATE),
}
W, H = 13.333, 7.5


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.parts: list[tuple[str, str]] = []
        self.section = 0

    # ------------------------------------------------------------ primitives
    def blank(self, bg=PAPER):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.rect(s, 0, 0, W, H, bg)
        return s

    @staticmethod
    def rect(s, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None):
        sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line
            sh.line.width = Pt(1)
        sh.shadow.inherit = False
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = 0.08
        return sh

    @staticmethod
    def text(s, x, y, w, h, runs, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, spacing=4, shape=None):
        """runs: str, or list of paragraphs; a paragraph is str or list of (text, overrides)."""
        if shape is None:
            shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for m in ("margin_left", "margin_right"):
            setattr(tf, m, Inches(0.08))
        tf.margin_top = tf.margin_bottom = Inches(0.04)
        paras = [runs] if isinstance(runs, str) else runs
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            if i:
                p.space_before = Pt(spacing)
            pieces = [(para, {})] if isinstance(para, str) else para
            for txt, ov in pieces:
                r = p.add_run()
                r.text = txt
                f = r.font
                f.name = "Calibri"
                f.size = Pt(ov.get("size", size))
                f.bold = ov.get("bold", bold)
                f.italic = ov.get("italic", False)
                f.color.rgb = ov.get("color", color)
        return shape

    def picture(self, s, path: Path, x, y, w, h):
        with Image.open(path) as im:
            ar = im.width / im.height
        pw, ph = w, w / ar
        if ph > h:
            ph, pw = h, h * ar
        s.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                             Inches(pw), Inches(ph))

    def notes(self, s, title, say):
        s.notes_slide.notes_text_frame.text = say
        self.parts.append((title, say))

    # ------------------------------------------------------------ chrome
    def chrome(self, s, title, status=None):
        name, col = SECTIONS[self.section]
        self.rect(s, 0, 0, 0.16, H, col)
        self.text(s, 0.55, 0.32, 9, 0.35, name.upper(), size=12, bold=True, color=col)
        self.text(s, 0.55, 0.6, 9.6 if status else 12.2, 0.8, title, size=28, bold=True, color=NAVY)
        if status:
            label, scol = STATUS[status]
            pill = self.rect(s, W - 3.05, 0.42, 2.55, 0.42, scol, MSO_SHAPE.ROUNDED_RECTANGLE)
            pill.adjustments[0] = 0.5
            self.text(s, 0, 0, 0, 0, label, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE, shape=pill)
        # footer: progress through the sections + slide number
        n = len(self.prs.slides)
        seg = 1.1
        for i, (_, c) in enumerate(SECTIONS):
            self.rect(s, 0.55 + i * (seg + 0.08), H - 0.36, seg, 0.07, c if i <= self.section else LINE)
        self.text(s, 6.6, H - 0.5, 4.5, 0.3, "dxFeed demo assessment · Irakli", size=10, color=SLATE)
        self.text(s, W - 1.3, H - 0.5, 0.8, 0.3, str(n), size=10, color=SLATE, align=PP_ALIGN.RIGHT)

    def banner(self, s, msg, y=6.18, color=NAVY, x=0.55, w=W - 1.05):
        b = self.rect(s, x, y, w, 0.72, color, MSO_SHAPE.ROUNDED_RECTANGLE)
        self.text(s, 0, 0, 0, 0, msg, size=17, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, shape=b)
        b.text_frame.margin_left = Inches(0.25)

    def points(self, s, x, y, w, items, color=TEAL, size=16, gap=0.12):
        """Short numbered points with coloured markers. items: list of (bold lead, rest)."""
        cy = y
        for i, (lead, rest) in enumerate(items, 1):
            dot = self.rect(s, x, cy + 0.04, 0.42, 0.42, color, MSO_SHAPE.OVAL)
            self.text(s, 0, 0, 0, 0, str(i), size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE, shape=dot)
            lines = max(1, round(len(lead + rest) / (w * 7.2 * 16 / size)) + 1)
            hgt = 0.33 * lines * size / 16
            self.text(s, x + 0.55, cy, w - 0.55, hgt,
                      [[(lead + " ", {"bold": True, "color": NAVY}), (rest, {})]], size=size)
            cy += hgt + gap
        return cy

    def card(self, s, x, y, w, h, head, body, color, size=14):
        self.rect(s, x, y, w, h, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, line=LINE)
        self.rect(s, x, y, w, 0.09, color)
        self.text(s, x + 0.15, y + 0.2, w - 0.3, 0.45, head, size=17, bold=True, color=color)
        self.text(s, x + 0.15, y + 0.7, w - 0.3, h - 0.8, body, size=size, color=INK, spacing=6)

    # ------------------------------------------------------------ layouts
    def chart_slide(self, title, status, image, items, takeaway, say, color=TEAL, img_w=8.3):
        s = self.blank()
        self.chrome(s, title, status)
        self.rect(s, 0.45, 1.5, img_w + 0.2, 4.5, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, line=LINE)
        self.picture(s, image, 0.55, 1.6, img_w, 4.3)
        x = img_w + 1.0
        self.points(s, x, 1.65, W - x - 0.45, items, color=color, size=15)
        self.banner(s, takeaway)
        self.notes(s, title, say)
        return s

    def wide_chart_slide(self, title, status, image, takeaway, say, side=None, color=TEAL):
        s = self.blank()
        self.chrome(s, title, status)
        self.rect(s, 0.45, 1.5, W - 0.9, 4.5, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, line=LINE)
        if side:
            self.picture(s, image, 0.55, 1.6, 9.0, 4.3)
            self.rect(s, 9.75, 1.75, 0.03, 4.0, LINE)
            self.points(s, 9.95, 1.75, 2.85, side, color=color, size=14)
        else:
            self.picture(s, image, 0.55, 1.6, W - 1.1, 4.3)
        self.banner(s, takeaway)
        self.notes(s, title, say)
        return s

    def divider(self, number, title, promise, minutes):
        self.section += 1
        name, col = SECTIONS[self.section]
        s = self.blank(NAVY)
        self.rect(s, 0, 0, 0.3, H, col)
        self.text(s, 0.9, 1.5, 4, 1.4, number, size=96, bold=True, color=col)
        self.text(s, 0.9, 3.05, 11, 0.9, title, size=44, bold=True, color=WHITE)
        self.text(s, 0.9, 4.0, 11, 0.9, promise, size=20, color=rgb("#C9D6E8"))
        for i, (_, c) in enumerate(SECTIONS[1:], 1):
            self.rect(s, 0.9 + (i - 1) * 1.55, 5.6, 1.45, 0.1, c if i <= self.section else rgb("#33415C"))
        self.text(s, 0.9, 5.85, 6, 0.4, f"about {minutes} minutes", size=14, color=rgb("#C9D6E8"))
        self.notes(s, f"Section: {title}", f"{title}. About {minutes} minutes. {promise}")


def kpis(d: Deck, s, tiles, y=1.65, h=2.05):
    w = (W - 1.1 - 3 * 0.25) / 4
    for i, (val, lab, col) in enumerate(tiles):
        x = 0.55 + (i % 4) * (w + 0.25)
        yy = y + (i // 4) * (h + 0.25)
        d.rect(s, x, yy, w, h, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, line=LINE)
        d.rect(s, x, yy, 0.09, h, col)
        d.text(s, x + 0.25, yy + 0.25, w - 0.4, 0.9, val, size=40, bold=True, color=col)
        d.text(s, x + 0.25, yy + 1.2, w - 0.4, 0.8, lab, size=14, color=INK)


def build() -> tuple[Path, list, int]:
    import deck_charts
    dch = deck_charts.main()
    rc = json.loads((EV / "recheck" / "recheck.json").read_text())
    total = rc["totals"]["total_candles"]
    d = Deck()

    # ================================================================ 1 title
    s = d.blank(NAVY)
    d.rect(s, 0, 0, 0.3, H, CORAL)
    d.rect(s, 0.3, 0, 0.1, H, AMBER)
    d.text(s, 0.95, 1.7, 11.5, 1.2, "How reliable is the dxFeed demo feed?", size=46, bold=True, color=WHITE)
    d.text(s, 0.95, 2.95, 11.5, 0.8, "Data consistency, behaviour under load, and what I would do about it",
           size=22, color=rgb("#C9D6E8"))
    d.text(s, 0.95, 4.3, 11.5, 0.5, "Irakli · L2 Support Engineer assessment · Task C", size=16, bold=True, color=WHITE)
    d.text(s, 0.95, 4.75, 11.5, 0.5, "Measured on Friday 16 October 2026 · every number reproducible offline from evidence/",
           size=14, color=rgb("#C9D6E8"))
    d.notes(s, "Title",
            "Hi, I'm Irakli. For Task C I treated the demo feed the way I would treat a client saying 'your data "
            "looks wrong'. I pulled enough history to test it properly, rebuilt the bigger candles from the smaller "
            "ones, compared everything, and then put some polite load on it. All of this was on Friday 16 October. "
            "Everything I show today can be recomputed from the saved files without touching the demo again.")

    # ================================================================ 2 agenda
    s = d.blank()
    d.chrome(s, "What I will cover")
    rows = [
        ("01", "Task A calendar", "Where the holiday data comes from, and what I could not get", "3 min", AMBER),
        ("02", "Data consistency", "Delay, the silent history cap, volume, prices, FX sessions", "15 min", TEAL),
        ("03", "Behaviour under load", "What happens when I add parallel requests", "9 min", CORAL),
        ("04", "Risk, incident, improvements", "What I would escalate, a drill, three proposals", "10 min", GREEN),
        ("", "Questions", "", "5 min", SLATE),
    ]
    for i, (num, name, what, mins, col) in enumerate(rows):
        y = 1.7 + i * 0.92
        d.rect(s, 0.55, y, W - 1.1, 0.78, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, line=LINE)
        d.rect(s, 0.55, y, 0.12, 0.78, col)
        d.text(s, 0.85, y + 0.12, 0.9, 0.55, num, size=24, bold=True, color=col, anchor=MSO_ANCHOR.MIDDLE)
        d.text(s, 1.75, y + 0.12, 4.0, 0.55, name, size=20, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        d.text(s, 5.6, y + 0.12, 6.0, 0.55, what, size=15, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
        d.text(s, 11.3, y + 0.12, 1.4, 0.55, mins, size=15, bold=True, color=col, align=PP_ALIGN.RIGHT,
               anchor=MSO_ANCHOR.MIDDLE)
    d.notes(s, "Agenda",
            "Four parts. Most of the time goes to data consistency, because that is where the interesting problems "
            "are. Every slide has a tag in the corner: verified means I reproduced it from the saved data, "
            "hypothesis means it fits but I could not prove it, not done means exactly that.")

    # ================================================================ 3 at a glance
    s = d.blank()
    d.chrome(s, "The short version", "verified")
    kpis(d, s, [
        (f"{total:,}", "candles pulled and checked", NAVY),
        ("900 s", "delay on every symbol: the feed is 15 min behind", TEAL),
        ("8,000", "candles per series, then it stops and still says OK", CORAL),
        ("130 + 1", "AAPL hours where 1h volume ≠ sum of 5m (ETH + RTH)", AMBER),
        ("50", "EUR/USD 1-minute candles outside the FX session", AMBER),
        ("~3 req/s", "throughput ceiling; more workers just wait longer", CORAL),
        ("0", "HTTP errors in 648 load requests, and no 429", GREEN),
        ("60", "holiday rows from the Task A scraper", BLUE),
    ])
    d.notes(s, "The short version",
            "If you remember one slide, this is it. The feed is delayed by exactly fifteen minutes. History silently "
            "stops at about eight thousand candles. Hourly volume for US stocks doesn't match the five-minute "
            "candles outside regular hours. FX prints a few candles when the market is closed. And under load the "
            "service queues instead of saying slow down. All latencies today are in seconds, not milliseconds.")

    # ================================================================ 4 how I worked
    s = d.blank()
    d.chrome(s, "How I worked")
    steps = [
        ("Pull", "96 chunked requests, 4 instruments, 1m / 5m / 1h", NAVY),
        ("Validate", "gaps, order, grid, NaN, sessions, holidays, DST", TEAL),
        ("Rebuild", "5m from 1m, 1h from 5m, compare candle by candle", GREEN),
        ("Compare", "SSE stream and the Radar web tool", BLUE),
        ("Load", "1 → 50 workers with automatic stop rules", AMBER),
        ("Recheck", "every headline number recomputed from disk", CORAL),
    ]
    bw = (W - 1.1 - 5 * 0.12) / 6
    for i, (name, what, col) in enumerate(steps):
        x = 0.55 + i * (bw + 0.12)
        sh = d.rect(s, x, 1.75, bw, 1.0, col, MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
        d.text(s, 0, 0, 0, 0, name, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE, shape=sh)
        sh.text_frame.margin_left = sh.text_frame.margin_right = Inches(0)
        d.text(s, x + 0.1, 2.9, bw - 0.2, 1.2, what, size=13, color=INK, align=PP_ALIGN.CENTER)
    d.card(s, 0.55, 3.85, 3.95, 2.15, "Instruments", [
        "AAPL · US stock", "TSLA · high volatility (no crypto on demo)", "EUR/USD · FX, bid candles only",
        "/ESZ26 · futures, used as a control"], NAVY, size=13)
    d.card(s, 4.69, 3.85, 3.95, 2.15, "Logged for every call", [
        "URL and parameters, UTC time", "latency, status, body hash", "raw body saved (gzip)"], TEAL, size=13)
    d.card(s, 8.83, 3.85, 3.95, 2.15, "Polite by design", [
        "one client, think time 0.25 s", "stop on 429, errors, p50 > 8 s", "budget of 3,500 requests"], GREEN, size=13)
    d.banner(s, "Rule I kept: a number goes on a slide only if I can recompute it offline.", y=6.2)
    d.notes(s, "How I worked",
            "This is the pipeline. I pulled history in chunks, validated each series, rebuilt the bigger timeframes "
            "from the smaller ones, cross-checked against the stream and Radar, then ran the load test. At the end "
            "I re-ran every headline number offline, and that re-check corrected four of my own earlier claims. "
            "I'll point them out when we get there. Crypto was in the brief, but every crypto symbol timed out "
            "on the demo, so TSLA is my high-volatility instrument.")

    # ================================================================ 01 Task A
    d.divider("01", "Task A: holiday calendar", "The calendar Task C uses to tell a holiday from a gap", 3)
    d.wide_chart_slide(
        "NYSE and NASDAQ agree on every 2026 date", "partial", TA / "calendar_2026.png",
        "60 rows, weekly on GitHub Actions. CME is not scraped because its terms forbid it.",
        "I built a small scraper with one adapter per exchange. It writes one CSV per exchange plus a combined "
        "file, and validates them. NYSE and NASDAQ match on every closed day and early close in 2026. The "
        "overall result is WARN, not FAIL. CME is a deliberate gap: its site terms forbid scraping, so the adapter "
        "only reads files dropped in by hand. OPRA has only published the 2025 PDF; the 2026 one returned 403. "
        "I uploaded the CSVs to Drive myself. The automated Drive upload is written but has not run with real "
        "credentials.",
        side=[
            ("60 rows:", "NYSE 34, NASDAQ 12, OPRA 14."),
            ("NASDAQ", "publishes the current year only."),
            ("OPRA", "2026 PDF not out yet (403)."),
            ("Drive:", "uploaded by hand; API path untested."),
        ],
        color=AMBER,
    )

    # ================================================================ 02 data
    d.divider("02", "Data consistency", "What is right, and the problems that do not raise an error", 15)
    d.chart_slide(
        "What is right: the basics are clean", "verified", EV / "plots" / "pass1" / "per_day_1m_counts.png",
        [
            ("0", "timestamps off the grid, out of order, or NaN."),
            ("Holidays", "match my Task A calendar exactly."),
            ("21 duplicates", "at chunk edges, all identical."),
            ("DST", "handled; first bar is 00:00 ET."),
        ],
        f"{total:,} candles: the price path is sound. The problems are quieter than that.",
        "I want to say what works before the problems. Across all twelve series there are no misaligned "
        "timestamps, nothing out of order, no NaN prices. Equities are empty on every NYSE holiday, matching "
        "Task A, and EUR/USD trades through them, which is correct. The 21 duplicates come from toTime being "
        "inclusive, and they carry identical values, so dropping them is safe.",
        color=GREEN,
    )
    d.chart_slide(
        "Every quote is exactly 15 minutes old", "verified", RCH / "quote_age.png",
        [
            ("Median 900.6 s", "on REST, SSE and Radar."),
            ("Load", "did not change it. Latency moved, freshness did not."),
            ("Nothing", "in the payload says the data is delayed."),
        ],
        "Any freshness alert has to start from a 900-second floor, not from zero.",
        "Quote age sits on nine hundred seconds for every symbol, on every surface. That's fine for a demo, but "
        "nothing in the response says it is delayed. A client who builds a 'data is stale' alert from zero will "
        "page all day. I'd want the delay in the payload.",
    )
    d.chart_slide(
        "History stops at 8,000 candles and still says OK", "verified", RCH / "candle_totals.png",
        [
            ("All four", "1-minute series stop at 8,000 or 8,001."),
            ("Older window:", "status OK, empty body, 15 bytes."),
            ("So it is retention,", "not a response-size limit."),
        ],
        "A two-week backtest silently loses a week. To the client it looks like a holiday.",
        "This is the one I'd escalate first. I asked for two weeks of one-minute AAPL and got the newest eight "
        "thousand, with status OK. The decisive test was a window completely before that edge: OK again, empty "
        "body. It may well be an intentional demo quota, and that's fine. The problem is the silence. A "
        "truncated flag, or one sentence in the docs, fixes it. I can show this live with chunks.jsonl.",
        color=CORAL,
    )
    d.wide_chart_slide(
        "Hourly volume breaks outside regular hours",
        "partial", RCH / "aapl_volume_heatmap.png",
        "AAPL: 130 extended-hours mismatches and 1 in regular hours. TSLA: 118 + 2. Cause not proven.",
        "This is the most interesting finding. Each cell is one hour; red means the served hourly bar has more "
        "volume than its five-minute bars add up to. The boxed regular-hours block is white, and outside it the "
        "difference goes up to 99 percent. There is one blue cell inside regular hours, 8 October 11:00 ET, "
        "plus 2 percent the other way. TSLA has two exceptions on the same day and hours, which points to a "
        "market-wide event, not an AAPL thing. Re-fetching 11 minutes later gave the same values. The test "
        "that would explain it needs TimeAndSale, and that timed out on the demo. So: real pattern, cause open.",
        side=[
            ("Red", "= hourly bar has more volume, up to 99 %."),
            ("Boxed", "regular hours are clean except one cell."),
            ("TSLA", "shows the same shape."),
            ("1m and 5m", "agree on every complete bucket."),
        ],
        color=CORAL,
    )
    d.chart_slide(
        "Which volume is right? Not decided yet", "partial", RCH / "volume_reconcile.png",
        [
            ("14 Oct daily candle:", "hourly −0.13 %, 1-minute −0.15 %."),
            ("Hourly", "is never below 1-minute on any day."),
            ("One day", "is not enough to pick a reference."),
        ],
        "What I would tell a client: we are confirming which series is the reference. No workaround yet.",
        "This slide changed my advice. My first draft said 'build extended-hours analytics from one-minute "
        "data'. The data doesn't support that: on the only day with a daily candle, the hourly sum is slightly "
        "closer. One data point, so I don't pick a winner. tho=true might be a workaround but I never tested it.",
        color=AMBER,
    )
    d.chart_slide(
        "Prices are almost, not perfectly, consistent", "corrected", RCH / "ohlc_mismatches.png",
        [
            ("6 hourly candles", "differ from their 5-minute rebuild."),
            ("3", "are the bar still forming at fetch time."),
            ("3", "are complete hours. Largest gap 0.023 %."),
        ],
        "I first wrote 'OHLC always agrees'. It doesn't. The gaps are tiny, but real.",
        "I'm correcting myself here. An earlier draft said prices always agree. They nearly do. Three of the six "
        "differences are the bar that was still forming, which is expected. Three are complete hours, and two "
        "of those are TSLA on 8 October, the same block as the volume exception. Small, but I'd rather say it.",
        color=AMBER,
    )
    d.chart_slide(
        "EUR/USD prints 50 candles while FX is closed", "corrected", RCH / "fx_outside_session.png",
        [
            ("7", "after the Friday close, 0 Saturday, 43 Sunday early."),
            ("170", "Sunday-evening candles are the normal reopen."),
            ("ES futures:", "0 outside session, so the method works."),
        ],
        "I first counted 214 by UTC weekday. Against the real FX session it is 50.",
        "FX trades Sunday 17:00 to Friday 17:00 New York time. My first count used the UTC weekday and got 214, "
        "but 170 of those are just the Sunday reopen. Against the real session it's 50. They might be indicative "
        "quotes or a heartbeat; I set up a weekend logger to tell them apart, but it stopped on Saturday, so that "
        "stays open. The client risk is session logic: 'last candle is fresh' would call FX open on Sunday morning.",
        color=AMBER,
    )

    # edge cases table
    s = d.blank()
    d.chrome(s, "Small things that turn into tickets", "verified")
    rows = [
        ("EUR/USD{=1m}", "OK, no candles. Only price=bid works", "Silent empty result"),
        ("Daily candle", "Stamped 00:00 UTC = 20:00 ET the day before", "Wrong date once shown in ET"),
        ("toTime", "Inclusive: 60 minutes returns 61 candles", "Duplicates at chunk edges"),
        ("Unknown symbol", "HTTP 200 + TIMED_OUT after 3.3 s (29 s under load)", "Looks like success, slows everything"),
        ("Forming candle", "No flag on it (eventFlags = 0)", "Client stores a bar that changes"),
        ("SSE in delay window", "NaN candles dated 2038-01-19", "Parsers store garbage"),
        ("Radar vs REST", "Radar shows NVDA, REST times it out", "'But it works in Radar'"),
    ]
    tbl = s.shapes.add_table(len(rows) + 1, 3, Inches(0.55), Inches(1.6), Inches(W - 1.1), Inches(4.4)).table
    for j, (hd, wd) in enumerate([("Where", 2.9), ("What happens", 5.6), ("What the client sees", 3.73)]):
        tbl.columns[j].width = Inches(wd)
        c = tbl.cell(0, j)
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
        d.text(s, 0, 0, 0, 0, hd, size=14, bold=True, color=WHITE, shape=c)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = WHITE if i % 2 else MIST
            d.text(s, 0, 0, 0, 0, val, size=14, bold=(j == 0), color=CORAL if j == 2 else (NAVY if j == 0 else INK),
                   shape=c)
    d.banner(s, "None of these is an error. Each one cost me an experiment, so each will cost a client a ticket.")
    d.notes(s, "Edge cases",
            "These are the small ones. None of them throws an error, which is exactly why they turn into "
            "tickets. My favourite is the 2038 date: if you open the stream inside the delay window you get NaN "
            "candles dated 2038. A parser that ignores the flags stores them. Radar versus REST is just different "
            "symbol lists, but support will hear 'it works in Radar' a lot.")

    # ================================================================ 03 load
    d.divider("03", "Behaviour under load", "What happens when one client adds parallel requests", 9)
    d.wide_chart_slide(
        "More workers did not buy more throughput", "corrected", dch["load"],
        "Throughput stays near 3 req/s while median latency goes 0.5 s → 8.3 s. Seconds, not milliseconds.",
        "Market-hours run. Left: throughput measured through completion. It jumps from one to five workers and "
        "then stays flat around three requests a second, and drops a little at fifty. Right: latency. Median "
        "goes from half a second to eight seconds, p95 to twenty-six. I want to be careful about two things. An "
        "earlier version said throughput rose to 4.6; that number ignored requests still in flight, and it's "
        "wrong. And these are seconds. Zero errors, no 429, no Retry-After: the service queues, it never tells "
        "the client to back off.",
        side=[
            ("Pre-market", "plateau was lower: ~1.6 req/s."),
            ("648", "requests, 0 HTTP errors, no 429."),
            ("Earlier '4.6 req/s'", "ignored in-flight requests."),
        ],
        color=CORAL,
    )
    s = d.blank()
    d.chrome(s, "Candle requests queue first, and timeouts stretch", "verified")
    for i, (img, cap) in enumerate([(RCH / "latency_by_request_type.png", "Median latency by request type (seconds)"),
                                    (RCH / "timed_out_stretch.png", "TIMED_OUT response time (seconds)")]):
        x = 0.45 + i * 6.3
        d.rect(s, x, 1.5, 6.1, 4.0, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, line=LINE)
        d.picture(s, img, x + 0.1, 1.6, 5.9, 3.45)
        d.text(s, x + 0.1, 5.08, 5.9, 0.35, cap, size=12, color=SLATE, align=PP_ALIGN.CENTER)
    d.banner(s, "At 25 workers quotes take 2.6 s, candles 11–12 s. An unknown symbol goes from 3.3 s to 29 s.",
             y=5.75, color=CORAL)
    d.text(s, 0.55, 6.55, 12, 0.4, "Percentiles count completed requests only, so at 25 and 50 workers the real tail is worse.",
           size=12, color=SLATE)
    d.notes(s, "Latency by request type",
            "Queueing isn't uniform. Quotes stay fairly quick while both candle requests wait eleven or twelve "
            "seconds. That fits per-request cost or separate queues, not one simple line. On the right, the "
            "timeout for an unknown symbol: three seconds alone, twenty-nine under load, because queue time is "
            "added before the server's own timeout starts. A dashboard mixing quotes and candles will look half "
            "alive.")

    s = d.blank()
    d.chrome(s, "Was the bottleneck my client? No.", "verified")
    big = [("in-flight = setting", "Little's law at 5 and 10 workers: my workers really were all waiting on the server.", TEAL),
           ("52 ≥ 50", "Connection pool limit was above the highest concurrency, so it could never bind.", GREEN),
           ("1 client", "Per-IP throttle or shared capacity? I can't tell without a second IP. Not run.", AMBER)]
    for i, (num, txt, col) in enumerate(big):
        x = 0.55 + i * 4.15
        d.rect(s, x, 1.65, 3.95, 3.6, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, line=LINE)
        d.rect(s, x, 1.65, 3.95, 0.09, col)
        d.text(s, x + 0.2, 1.95, 3.6, 1.1, num, size=30 if len(num) > 8 else 44, bold=True, color=col)
        d.text(s, x + 0.2, 3.2, 3.6, 1.9, txt, size=16, color=INK)
    d.banner(s, "This shows how this shared demo behaves for one client. It is not a production capacity number.")
    d.notes(s, "Bottleneck",
            "The first question I asked myself was whether I was the bottleneck. Little's law says no: at five and "
            "ten workers the average in flight equals the setting, so they were all waiting on the server. The "
            "pool allowed 52 connections and I never ran more than 50. What I can't say is whether the cap is per "
            "IP or shared with other demo users. That needs a second client at the same time, and I didn't run "
            "it. If anyone asks for a production number: I won't extrapolate one from a free demo.")

    # ================================================================ 04 risk
    d.divider("04", "Risk, incident and improvements", "What I would escalate, how I'd run the incident, what I'd fix", 10)
    s = d.blank()
    d.chrome(s, "What I would escalate, in order")
    d.rect(s, 0.45, 1.5, 6.0, 4.55, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, line=LINE)
    d.picture(s, dch["risk"], 0.55, 1.6, 5.8, 4.35)
    for i, (n, label, lk, im) in enumerate(deck_charts.RISKS):
        col = CORAL if lk + im > 5 else AMBER if lk + im > 4 else TEAL
        y = 1.6 + i * 0.62
        dot = d.rect(s, 6.8, y, 0.46, 0.46, col, MSO_SHAPE.OVAL)
        d.text(s, 0, 0, 0, 0, str(n), size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE, shape=dot)
        d.text(s, 7.4, y + 0.02, 5.4, 0.45, label, size=17, color=NAVY, bold=n <= 2, anchor=MSO_ANCHOR.MIDDLE)
    d.banner(s, "Silent beats loud: a wrong answer with status OK is worse than an error.")
    d.notes(s, "Risks",
            "I ranked risks by how often I actually saw them and how much damage they do if nobody notices. "
            "Silent ones go to the top: the 8,000 cap, the volume difference, and staleness that looks like a "
            "closed market. Queueing without a 429 is next. The FX prints are real but low impact if clients use "
            "a proper session calendar.")

    s = d.blank()
    d.chrome(s, "Incident drill: 'AAPL hourly volume is corrupted'", "simulation")
    d.rect(s, 0.45, 1.45, W - 0.9, 2.55, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, line=LINE)
    d.picture(s, dch["timeline"], 0.55, 1.5, W - 1.1, 2.45)
    d.card(s, 0.55, 4.2, 6.0, 2.75, "First 15 minutes", [
        "Reproduce from public endpoints: −98.4 % on 2 Oct 18:00 ET.",
        "Count: 130 extended-hours, 1 regular-hours.",
        "Confirm 5m = 1m; check the daily total.",
        "Escalate with the CSV and one question: which series is the reference?"], TEAL, size=14)
    d.card(s, 6.78, 4.2, 6.0, 2.75, "What the client hears", [
        [("“We reproduced it. It affects hourly volume for US stocks, almost only outside regular hours. "
          "Prices agree. We are confirming the reference series. Next update 10:00 UTC.”", {"italic": True})],
        [("No root cause, no untested workaround.", {"bold": True, "color": CORAL})]], NAVY, size=14)
    d.notes(s, "Incident drill",
            "The scenario is invented, the numbers are real. A quant client's nightly reconciliation flags the "
            "18:00 hour. In the first fifteen minutes I reproduce it from public endpoints, count how many hours "
            "are affected, check five-minute against one-minute, and escalate with the CSV and one clear question. "
            "The client message gives facts and a next-update time. It doesn't guess a cause and it doesn't offer "
            "tho=true, because I never tested it.")

    s = d.blank()
    d.chrome(s, "Three things I would change", "proposal")
    cols = [
        ("Technical", TEAL, "Say what the response is",
         ["Add truncated, serverTime, delaySeconds to the response.", "Flag the forming candle.",
          "Return UNKNOWN_SYMBOL at once, not TIMED_OUT."], "Medium", "Removes the top two 'wrong data' tickets"),
        ("Monitoring", AMBER, "Catch it before the client does",
         ["One synthetic probe per asset class.", "Age vs calendar, 1h vs 1m volume, p95 at low concurrency.",
          "Budget the probes: the service queues."], "Small–medium", "Client-found incidents become ours"),
        ("Documentation", GREEN, "One 'demo facts' page",
         ["Delay, retention cap, symbol list.", "FX needs price=bid; toTime is inclusive.",
          "Radar vs REST differences."], "Small", "Fewer first-contact tickets"),
    ]
    for i, (area, col, head, body, cx, impact) in enumerate(cols):
        x = 0.55 + i * 4.15
        d.rect(s, x, 1.55, 3.95, 5.4, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, line=LINE)
        hb = d.rect(s, x, 1.55, 3.95, 0.6, col)
        d.text(s, 0, 0, 0, 0, area.upper(), size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, shape=hb)
        hb.text_frame.margin_left = Inches(0.2)
        d.text(s, x + 0.2, 2.3, 3.6, 0.6, head, size=19, bold=True, color=NAVY)
        d.text(s, x + 0.2, 2.95, 3.6, 2.3, ["• " + b for b in body], size=14, spacing=8)
        d.rect(s, x + 0.2, 5.4, 3.55, 0.03, LINE)
        d.text(s, x + 0.2, 5.5, 3.6, 0.4, [[("Effort  ", {"color": SLATE}), (cx, {"bold": True, "color": col})]], size=14)
        d.text(s, x + 0.2, 5.9, 3.6, 0.9, [[("Impact  ", {"color": SLATE}), (impact, {"bold": True, "color": NAVY})]], size=14)
    d.notes(s, "Improvements",
            "Three proposals. The technical one matters most, because a truncated flag fixes the silent cap with "
            "an additive JSON field. Monitoring: the scripts I wrote are most of a probe already; what's missing "
            "is wiring and a session calendar. Documentation is the cheapest: one page, and my report is the draft.")

    # corrections
    s = d.blank()
    d.chrome(s, "Four things I got wrong first, and fixed", "corrected")
    fixes = [
        ("“214 weekend FX candles”", "50 outside the real FX session", "I counted by UTC weekday"),
        ("“Throughput rises to 4.6 req/s”", "Flat at ~3 req/s", "I ignored requests still in flight"),
        ("“OHLC always agrees”", "6 small differences, ≤ 0.023 %", "My first check was too loose"),
        ("“Rebuild ETH volume from 1m”", "Reference series not decided", "The daily candle says otherwise"),
    ]
    for i, (old, new, why) in enumerate(fixes):
        y = 1.6 + i * 1.12
        d.rect(s, 0.55, y, 4.6, 0.92, rgb("#FDECE7"), MSO_SHAPE.ROUNDED_RECTANGLE)
        d.text(s, 0.75, y + 0.1, 4.3, 0.75, old, size=16, color=CORAL, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        arrow = d.rect(s, 5.3, y + 0.26, 0.6, 0.4, SLATE, MSO_SHAPE.RIGHT_ARROW)
        d.rect(s, 6.05, y, 3.6, 0.92, rgb("#E6F4F1"), MSO_SHAPE.ROUNDED_RECTANGLE)
        d.text(s, 6.25, y + 0.1, 3.3, 0.75, new, size=16, color=GREEN, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        d.text(s, 9.85, y + 0.1, 3.0, 0.75, why, size=14, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
        _ = arrow
    d.banner(s, "Re-running everything offline is what caught these. That re-check is now part of the method.",
             y=6.18, color=GREEN)
    d.notes(s, "Corrections",
            "I'd rather show you this than have you find it. Each of these was in an earlier draft. The offline "
            "re-check caught all four, and every slide today uses the corrected numbers.")

    # open items
    s = d.blank()
    d.chrome(s, "What is still open", "not done")
    items = [
        ("Waiting on the clock", CORAL, ["Pass 2 re-pull after Monday 19 Oct 15:30Z",
                                         "Weekend logger: stopped Sat 14:28Z, Sunday not covered"]),
        ("Blocked on the demo", AMBER, ["TimeAndSale / Trade history: TIMED_OUT or empty",
                                        "So the volume cause stays a hypothesis"]),
        ("Not run", SLATE, ["SPY / MSFT check of 8 Oct 11:00 ET; tho=true",
                            "Second client / IP; burst and batching tests"]),
        ("Task A", BLUE, ["CME: manual files only (terms)", "OPRA 2026 PDF not published; Drive API untested"]),
    ]
    for i, (head, col, body) in enumerate(items):
        x = 0.55 + (i % 2) * 6.2
        y = 1.6 + (i // 2) * 2.25
        d.card(s, x, y, 6.0, 2.05, head, ["• " + b for b in body], col, size=15)
    d.notes(s, "Open items",
            "If you ask what I didn't verify, this is the answer. Pass 2 and the logger are only waiting on the "
            "clock. The volume cause is blocked because the demo doesn't serve TimeAndSale history. The rest I "
            "simply didn't run, and I'm not going to present them as results.")

    # questions
    s = d.blank(NAVY)
    d.rect(s, 0, 0, 0.3, H, GREEN)
    d.text(s, 0.95, 1.6, 11, 1.2, "Questions", size=54, bold=True, color=WHITE)
    d.text(s, 0.95, 2.8, 11, 0.6, "Happy to open any evidence file live.", size=22, color=rgb("#C9D6E8"))
    d.card(s, 0.95, 3.8, 5.6, 2.7, "I can show live", [
        "The 8,000 cap: a long fromTime, or chunks.jsonl",
        "The AAPL −99 % and +2 % volume rows in the CSV",
        "Radar screenshots and the SSE comparison"], TEAL, size=15)
    d.card(s, 6.8, 3.8, 5.6, 2.7, "I won't claim", [
        "A root cause for the volume split",
        "That the FX prints are trades",
        "A production capacity number"], CORAL, size=15)
    d.notes(s, "Questions",
            "That's it from me. If you ask whether the 8,000 cap is just a quota: possibly, and the missing "
            "signal is still the defect. If you ask about production capacity: the monitoring design transfers, "
            "the numbers don't.")

    d.prs.save(OUT)
    return OUT, d.parts, len(d.prs.slides)


if __name__ == "__main__":
    path, parts, n = build()
    print(path, n)
