"""Shared look and feel for the Word reports and charts of Tasks A, B and C.

One palette, one matplotlib style and a handful of python-docx helpers (cover page, callout
boxes, shaded tables, figures with captions) so the three reports read as one package.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402,F401
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

NAVY = "#14213D"
TEAL = "#0F8B8D"
GREEN = "#2A9D8F"
AMBER = "#E9A23B"
CORAL = "#E76F51"
BLUE = "#3A86FF"
PURPLE = "#7B2CBF"
SLATE = "#5C677D"
LIGHT = "#F4F6FA"
INK = "#1B1B1E"

STATUS_COLORS = {"verified": GREEN, "done": GREEN, "hypothesis": AMBER, "partial": AMBER,
                 "not done": CORAL, "blocked": CORAL}

AUTHOR = "Irakli — L2 Support Engineer candidate"


def mpl_style():
    plt.rcParams.update({
        "font.family": "DejaVu Sans",
        "font.size": 11,
        "axes.titlesize": 14,
        "axes.titleweight": "bold",
        "axes.titlecolor": NAVY,
        "axes.labelcolor": INK,
        "axes.edgecolor": "#C9CED6",
        "axes.spines.top": False,
        "axes.spines.right": False,
        "axes.grid": True,
        "axes.axisbelow": True,
        "grid.color": "#E6E9EF",
        "grid.linewidth": 0.8,
        "xtick.color": SLATE,
        "ytick.color": SLATE,
        "legend.frameon": False,
        "figure.facecolor": "white",
        "savefig.bbox": "tight",
        "savefig.dpi": 170,
    })


def _rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _shade(cell, hex_color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tc_pr.append(shd)


def _cell_margins(cell, cm=0.25):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(cm * 567)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


TEXT_WIDTH_CM = 16.6


def _set_widths(t, widths_cm):
    t.autofit = False
    tbl_pr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tbl_pr.append(lay)
    for ci, w in enumerate(widths_cm):
        t.columns[ci].width = Cm(w)
        for row in t.rows:
            row.cells[ci].width = Cm(w)


def _no_split(row, header=False):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    tr_pr.append(el)
    if header:
        tr_pr.append(OxmlElement("w:tblHeader"))


def new_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.12
    for name, size, color in (("Heading 1", 18, NAVY), ("Heading 2", 14, TEAL), ("Heading 3", 12, SLATE)):
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = _rgb(color)
        h.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        h.paragraph_format.space_after = Pt(6)
    return doc


def cover(doc: Document, title: str, subtitle: str, task_label: str, facts: list[tuple[str, str]]):
    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = band.rows[0].cells[0]
    _shade(cell, NAVY)
    _cell_margins(cell, 0.6)
    p = cell.paragraphs[0]
    r = p.add_run(task_label.upper())
    r.font.size, r.font.bold, r.font.color.rgb = Pt(11), True, _rgb(AMBER)
    p2 = cell.add_paragraph()
    r = p2.add_run(title)
    r.font.size, r.font.bold, r.font.color.rgb = Pt(28), True, _rgb("#FFFFFF")
    p3 = cell.add_paragraph()
    r = p3.add_run(subtitle)
    r.font.size, r.font.color.rgb = Pt(13), _rgb("#DDE3EC")
    doc.add_paragraph()
    t = doc.add_table(rows=len(facts), cols=2)
    t.style = "Table Grid"
    for i, (k, v) in enumerate(facts):
        a, b = t.rows[i].cells
        _shade(a, LIGHT)
        a.text, b.text = k, v
        a.paragraphs[0].runs[0].font.bold = True
    _set_widths(t, [4.5, TEXT_WIDTH_CM - 4.5])
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(f"{AUTHOR} · {dt.date(2026, 10, 18):%d %B %Y}")
    r.font.color.rgb = _rgb(SLATE)
    r.italic = True
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def para(doc: Document, text: str, bold_lead: str | None = None, italic=False):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead + " ")
        r.bold = True
    r = p.add_run(text)
    r.italic = italic
    return p


def bullets(doc: Document, items, style="List Bullet"):
    for it in items:
        if isinstance(it, tuple):
            p = doc.add_paragraph(style=style)
            r = p.add_run(it[0] + " ")
            r.bold = True
            p.add_run(it[1])
        else:
            doc.add_paragraph(it, style=style)


def numbered(doc: Document, items):
    for n, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
        r = p.add_run(f"{n}.\t")
        r.bold, r.font.color.rgb = True, _rgb(TEAL)
        if isinstance(it, tuple):
            r = p.add_run(it[0] + " ")
            r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)


def callout(doc: Document, title: str, text: str | list[str], color: str = TEAL):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    _shade(c, color)
    _cell_margins(c, 0.35)
    _set_widths(t, [TEXT_WIDTH_CM])
    _no_split(t.rows[0])
    p = c.paragraphs[0]
    r = p.add_run(title)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(12), _rgb("#FFFFFF")
    for line in ([text] if isinstance(text, str) else text):
        q = c.add_paragraph()
        rr = q.add_run(line)
        rr.font.color.rgb = _rgb("#FFFFFF")
        rr.font.size = Pt(10.5)
    doc.add_paragraph()


def table(doc: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float] | None = None,
          header_color: str = NAVY, status_col: int | None = None, font_size: float = 9.5):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        _shade(c, header_color)
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(font_size), _rgb("#FFFFFF")
    for ri, row in enumerate(rows, start=1):
        for ci, v in enumerate(row):
            c = t.rows[ri].cells[ci]
            c.text = ""
            r = c.paragraphs[0].add_run(str(v))
            r.font.size = Pt(font_size)
            if status_col is not None and ci == status_col:
                col = STATUS_COLORS.get(str(v).lower())
                if col:
                    _shade(c, col)
                    r.font.color.rgb = _rgb("#FFFFFF")
                    r.bold = True
            elif ri % 2 == 0:
                _shade(c, LIGHT)
    if not widths_cm:
        widths_cm = [TEXT_WIDTH_CM / len(headers)] * len(headers)
    _set_widths(t, widths_cm)
    for i, row in enumerate(t.rows):
        _no_split(row, header=i == 0)
        if i < len(t.rows) - 1:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True
    doc.add_paragraph()
    return t


_fig_no = {"n": 0}


def figure(doc: Document, path: Path, caption: str, width_cm: float = 16.0):
    if not Path(path).exists():
        para(doc, f"[missing figure: {path}]", italic=True)
        return
    _fig_no["n"] += 1
    doc.add_picture(str(path), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Figure {_fig_no['n']}. ")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(9.5), _rgb(TEAL)
    r = p.add_run(caption)
    r.italic, r.font.size, r.font.color.rgb = True, Pt(9.5), _rgb(SLATE)


def reset_figures():
    _fig_no["n"] = 0




def footer(doc: Document, text: str):
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.text = text
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.size = Pt(8.5)
            r.font.color.rgb = _rgb(SLATE)


def kpi_strip(doc: Document, items: list[tuple[str, str, str]]):
    """A row of coloured tiles: (big value, label, colour)."""
    t = doc.add_table(rows=1, cols=len(items))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_widths(t, [TEXT_WIDTH_CM / len(items)] * len(items))
    for i, (val, label, color) in enumerate(items):
        c = t.rows[0].cells[i]
        _shade(c, color)
        _cell_margins(c, 0.25)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(val)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(18), _rgb("#FFFFFF")
        q = c.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = q.add_run(label)
        rr.font.size, rr.font.color.rgb = Pt(8.5), _rgb("#FFFFFF")
    doc.add_paragraph()
